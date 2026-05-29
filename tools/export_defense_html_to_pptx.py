#!/usr/bin/env python3
"""
Exporta index.html (?export=1) a PPTX con una imagen por diapositiva (fidelidad visual).

dom-to-pptx degrada tablas/diagramas y no incrusta imágenes locales (file://).
Este pipeline es el entregable fiable para *_html.pptx.

Requisitos:
    python -m pip install python-pptx pillow playwright
    playwright install chromium

Uso (raíz del repo):
    python tools/export_defense_html_to_pptx.py
"""

from __future__ import annotations

import asyncio
import shutil
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from defense_slides_manifest import OUTPUT_PPTX_FROM_HTML, SLIDES


def repo_root() -> Path:
    return Path(__file__).resolve().parents[1]


def html_path() -> Path:
    return repo_root() / "docs" / "presentacion-defensa-pfc" / "index.html"


def out_pptx() -> Path:
    return repo_root() / "docs" / OUTPUT_PPTX_FROM_HTML


async def capture_slide_pngs(tmp_dir: Path, viewport: tuple[int, int]) -> list[str]:
    from playwright.async_api import async_playwright

    html = html_path().resolve()
    url = html.as_uri() + "?export=1"
    vw, vh = viewport

    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": vw, "height": vh})
        await page.goto(url, wait_until="networkidle")
        await page.wait_for_function(
            "window.__PFC_EXPORT_READY__ === true", timeout=60_000
        )
        await page.evaluate(
            """async () => {
              await Promise.all(
                Array.from(document.images).map(
                  (img) =>
                    new Promise((resolve) => {
                      if (img.complete) { resolve(); return; }
                      const done = () => resolve();
                      img.addEventListener("load", done, { once: true });
                      img.addEventListener("error", done, { once: true });
                      setTimeout(done, 5000);
                    })
                )
              );
            }"""
        )
        ids: list[str] = await page.evaluate(
            """() => [...document.querySelectorAll('.slide')].map(s => s.dataset.id)"""
        )
        if len(ids) != len(SLIDES):
            print(
                f"AVISO: HTML tiene {len(ids)} slides, manifiesto {len(SLIDES)}",
                file=sys.stderr,
            )
        order: list[str] = []
        for sid in ids:
            await page.evaluate(
                f'document.querySelector("[data-id=\\"{sid}\\"]")'
                ".scrollIntoView({ block: 'start', behavior: 'instant' })"
            )
            await page.wait_for_timeout(350)
            path = tmp_dir / f"{sid}.png"
            await page.screenshot(path=str(path), full_page=False)
            order.append(sid)
            print(f"  captura {sid}")
        await browser.close()
    return order


def build_pptx_from_pngs(png_dir: Path, slide_ids: list[str]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_idx = 6 if len(prs.slide_layouts) > 6 else 5

    for sid in slide_ids:
        png = png_dir / f"{sid}.png"
        if not png.is_file():
            raise FileNotFoundError(f"Falta captura: {png}")
        slide = prs.slides.add_slide(prs.slide_layouts[blank_idx])
        slide.shapes.add_picture(
            str(png),
            Inches(0),
            Inches(0),
            width=prs.slide_width,
            height=prs.slide_height,
        )

    out = out_pptx()
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"Guardado: {out} ({len(slide_ids)} diapositivas, imagen full-bleed)")


def main() -> None:
    if not html_path().is_file():
        print(f"No encontrado: {html_path()}", file=sys.stderr)
        print("Ejecuta: python tools/generate_defense_html.py", file=sys.stderr)
        raise SystemExit(1)

    viewport = (1920, 1080)
    tmp = Path(tempfile.mkdtemp(prefix="pfc_html_export_"))
    try:
        print(f"Capturando HTML ({viewport[0]}x{viewport[1]})...")
        order = asyncio.run(capture_slide_pngs(tmp, viewport))
        print("Montando PPTX...")
        build_pptx_from_pngs(tmp, order)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
