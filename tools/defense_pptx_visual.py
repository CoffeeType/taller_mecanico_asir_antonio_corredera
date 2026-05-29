"""
Builders ppt-visual (Corporate Blue) para generate_defense_pptx.py.
Layouts inspirados en docs/presentacion-defensa-pfc/index.html:
diagramas centrados con object-fit contain, dual grid, números de slide.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from defense_slides_manifest import DefenseSlide

# Corporate Blue (ppt-visual) — estructura alineada al HTML (jerarquía + acento)
COLOR_PRIMARY = RGBColor(0x1E, 0x3A, 0x5F)
COLOR_SECONDARY = RGBColor(0x34, 0x98, 0xDB)
COLOR_ACCENT = RGBColor(0xE7, 0x4C, 0x3C)
COLOR_BG = RGBColor(0xF5, 0xF7, 0xFA)
COLOR_TEXT = RGBColor(0x2C, 0x3E, 0x50)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_ROW_ALT = RGBColor(0xE8, 0xEE, 0xF4)
COLOR_FRAME = RGBColor(0xE2, 0xE8, 0xF0)
COLOR_NUM_FADE = RGBColor(0x34, 0x98, 0xDB)

FONT_TITLE = "Montserrat"
FONT_BODY = "Open Sans"
FONT_FALLBACK = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_L = Inches(0.55)
CONTENT_TOP = Inches(1.05)
FOOTER_TOP = Inches(7.05)
TITLE_TOP = Inches(0.38)

# Equivalente HTML: max-height min(50vh, 400px) en 720p → ~3.65" de 7.5"
DIAGRAM_MAX_H = Inches(3.65)
DIAGRAM_MAX_W = Inches(11.8)
DUAL_MAX_H = Inches(3.55)
DUAL_COL_W = Inches(5.95)


def _set_font_safe(run, name: str, size_pt: int, bold: bool = False, color: RGBColor | None = None):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    try:
        run._element.rPr.rFonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii", name
        )
        run._element.rPr.rFonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi", name
        )
    except Exception:
        pass


def _fit_size(path: Path, max_w_in: float, max_h_in: float) -> tuple[float, float]:
    """Calcula ancho/alto en pulgadas manteniendo aspect ratio (contain)."""
    with Image.open(path) as im:
        w_px, h_px = im.size
    if w_px <= 0 or h_px <= 0:
        return max_w_in, max_h_in
    aspect = w_px / h_px
    box_aspect = max_w_in / max_h_in
    if aspect > box_aspect:
        w_in = max_w_in
        h_in = max_w_in / aspect
    else:
        h_in = max_h_in
        w_in = max_h_in * aspect
    return w_in, h_in


def set_notes(slide, text: str):
    notes = slide.notes_slide.notes_text_frame
    notes.clear()
    notes.text = text
    for p in notes.paragraphs:
        p.font.size = Pt(11)
        p.font.name = FONT_FALLBACK
        p.font.color.rgb = COLOR_TEXT


def _blank_slide(prs: Presentation):
    layout_idx = 6 if len(prs.slide_layouts) > 6 else 5
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG
    return slide


def _add_footer(slide):
    box = slide.shapes.add_textbox(MARGIN_L, FOOTER_TOP, Inches(12.2), Inches(0.35))
    box.text_frame.text = "Defensa PFC — Taller mecánico ASIR"
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    for r in p.runs:
        _set_font_safe(r, FONT_FALLBACK, 8, color=COLOR_TEXT)


def _add_slide_number(slide, index: int):
    """Número de sección estilo HTML (.slide-num)."""
    num = slide.shapes.add_textbox(Inches(0.5), Inches(0.12), Inches(2.5), Inches(0.55))
    num.text_frame.text = f"{index:02d}"
    for p in num.text_frame.paragraphs:
        for r in p.runs:
            _set_font_safe(r, FONT_TITLE, 36, bold=True, color=COLOR_NUM_FADE)
    try:
        for r in num.text_frame.paragraphs[0].runs:
            r.font.color.rgb = COLOR_NUM_FADE
    except Exception:
        pass


def _add_slide_title(slide, title: str, top=TITLE_TOP, size=24, left=Inches(2.35)):
    width = Inches(10.4)
    box = slide.shapes.add_textbox(left, top, width, Inches(0.62))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = title
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            _set_font_safe(r, FONT_TITLE, size, bold=True, color=COLOR_PRIMARY)


def _add_accent_bar(slide, left=Inches(0.35), top=CONTENT_TOP, height=Inches(5.65)):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.1), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_SECONDARY
    bar.line.fill.background()


def _fill_bullets(tf, bullets: list[str], size=15, compact=False):
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    gap = Pt(5) if compact else Pt(7)
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {line}" if not line.startswith("•") else line
        p.level = 0
        p.space_after = gap
        for r in p.runs:
            _set_font_safe(r, FONT_BODY, size, color=COLOR_TEXT)


def _add_image_frame(slide, left, top, width, height):
    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = COLOR_FRAME
    frame.line.color.rgb = COLOR_SECONDARY
    frame.line.width = Pt(1)
    return frame


def _place_image_contain(
    slide,
    image_path: Path,
    box_left: float,
    box_top: float,
    max_w_in: float,
    max_h_in: float,
    with_frame: bool = True,
) -> float:
    """Coloca imagen centrada en caja max_w x max_h (como HTML object-fit: contain). Devuelve borde inferior en pulgadas."""
    w_in, h_in = _fit_size(image_path, max_w_in, max_h_in)
    pad = 0.12
    frame_w = w_in + pad * 2
    frame_h = h_in + pad * 2
    cx = box_left + (max_w_in - frame_w) / 2
    cy = box_top + (max_h_in - frame_h) / 2
    if with_frame:
        _add_image_frame(slide, Inches(cx), Inches(cy), Inches(frame_w), Inches(frame_h))
    pic_left = cx + pad
    pic_top = cy + pad
    slide.shapes.add_picture(
        str(image_path), Inches(pic_left), Inches(pic_top), width=Inches(w_in), height=Inches(h_in)
    )
    return cy + frame_h


def render_title_slide(prs: Presentation, spec: DefenseSlide, slide_index: int):
    slide = _blank_slide(prs)
    band_h = Inches(3.35)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, band_h)
    band.fill.solid()
    band.fill.fore_color.rgb = COLOR_PRIMARY
    band.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.85), Inches(11.8), Inches(1.8))
    title_box.text_frame.word_wrap = True
    title_box.text_frame.text = spec.title
    for p in title_box.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            _set_font_safe(r, FONT_TITLE, 30, bold=True, color=COLOR_WHITE)

    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.15), Inches(3.5), Inches(11.0), Inches(3.4)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = COLOR_WHITE
    panel.line.color.rgb = COLOR_SECONDARY
    panel.line.width = Pt(1.5)

    sub_box = slide.shapes.add_textbox(Inches(1.45), Inches(3.7), Inches(10.4), Inches(2.95))
    _fill_bullets(sub_box.text_frame, spec.bullets, size=16)
    for p in sub_box.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER

    _add_footer(slide)
    set_notes(slide, spec.notes)
    return slide


def render_timeline_slide(prs: Presentation, spec: DefenseSlide, slide_index: int):
    slide = _blank_slide(prs)
    _add_slide_number(slide, slide_index)
    _add_slide_title(slide, spec.title)

    labels = [row[0] for row in spec.table_rows]
    n = len(labels)
    y_node = Inches(1.88)
    x_start = Inches(1.0)
    x_end = Inches(12.3)
    step = (x_end - x_start) / max(n - 1, 1)

    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, x_start, y_node + Inches(0.18), x_end - x_start, Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_SECONDARY
    line.line.fill.background()

    for i, label in enumerate(labels):
        cx = x_start + step * i
        oval = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, cx - Inches(0.2), y_node, Inches(0.4), Inches(0.4)
        )
        oval.fill.solid()
        oval.fill.fore_color.rgb = COLOR_PRIMARY if i % 2 == 0 else COLOR_SECONDARY
        oval.line.fill.background()
        lbl = slide.shapes.add_textbox(cx - Inches(0.42), y_node + Inches(0.48), Inches(0.84), Inches(0.32))
        lbl.text_frame.text = label
        for p in lbl.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                _set_font_safe(r, FONT_BODY, 10, bold=True, color=COLOR_PRIMARY)

    headers = spec.table_headers
    rows = spec.table_rows
    ncols = len(headers)
    nrows = 1 + len(rows)
    tbl_top = Inches(2.62)
    tbl = slide.shapes.add_table(
        nrows, ncols, MARGIN_L, tbl_top, Inches(12.2), Inches(0.3 * nrows + 0.12)
    ).table
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                _set_font_safe(r, FONT_BODY, 11, bold=True, color=COLOR_WHITE)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = val
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLOR_ROW_ALT
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    _set_font_safe(r, FONT_BODY, 10, color=COLOR_TEXT)

    _add_footer(slide)
    set_notes(slide, spec.notes)
    return slide


def render_content_accent(prs: Presentation, spec: DefenseSlide, slide_index: int):
    slide = _blank_slide(prs)
    _add_slide_number(slide, slide_index)
    _add_slide_title(slide, spec.title)
    _add_accent_bar(slide)
    box = slide.shapes.add_textbox(Inches(0.62), CONTENT_TOP, Inches(12.0), Inches(5.75))
    _fill_bullets(box.text_frame, spec.bullets)
    _add_footer(slide)
    set_notes(slide, spec.notes)
    return slide


def render_table_styled(prs: Presentation, spec: DefenseSlide, slide_index: int):
    slide = _blank_slide(prs)
    _add_slide_number(slide, slide_index)
    _add_slide_title(slide, spec.title)
    headers = spec.table_headers
    rows = spec.table_rows
    ncols = len(headers)
    nrows = 1 + len(rows)
    row_h = 0.36 if nrows > 4 else 0.42
    tbl = slide.shapes.add_table(
        nrows, ncols, MARGIN_L, CONTENT_TOP, Inches(12.2), Inches(row_h * nrows + 0.15)
    ).table
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                _set_font_safe(r, FONT_BODY, 12, bold=True, color=COLOR_WHITE)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = val
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLOR_ROW_ALT
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    _set_font_safe(r, FONT_BODY, 11, color=COLOR_TEXT)
    _add_footer(slide)
    set_notes(slide, spec.notes)
    return slide


def render_diagram_stack(
    prs: Presentation, spec: DefenseSlide, image_path: Path | None, slide_index: int
):
    """
    Como HTML .diagram-slide: título → figura centrada (max 50vh) → bullets debajo.
    """
    slide = _blank_slide(prs)
    _add_slide_number(slide, slide_index)
    _add_slide_title(slide, spec.title, size=22)

    img_top = 1.05
    box_left = 0.55
    max_w = 12.2
    max_h = 3.65

    bottom_y = img_top + max_h
    if image_path and image_path.is_file():
        bottom_y = _place_image_contain(
            slide, image_path, box_left, img_top, max_w, max_h, with_frame=True
        )
        bottom_y += 0.12

    bullet_top = max(bottom_y, 4.75)
    if bullet_top > 6.4:
        bullet_top = 4.75
    bullet_h = max(6.85 - bullet_top, 1.0)
    box = slide.shapes.add_textbox(MARGIN_L, Inches(bullet_top), Inches(12.2), Inches(bullet_h))
    _fill_bullets(box.text_frame, spec.bullets, size=14, compact=True)

    _add_footer(slide)
    set_notes(slide, spec.notes)
    return slide


def render_dual_images(
    prs: Presentation,
    spec: DefenseSlide,
    left: Path | None,
    right: Path | None,
    slide_index: int,
):
    """Como HTML .dual-images: dos columnas con contain y pies de figura."""
    slide = _blank_slide(prs)
    _add_slide_number(slide, slide_index)
    _add_slide_title(slide, spec.title, size=22)

    if spec.id == "data_model":
        captions = ("Fig. 2 — Modelo de datos", "Fig. 3 — Flujo de cita")
    else:
        captions = ("Grafana — dashboards", "docker ps — contenedores")

    img_top = 1.02
    max_h = 3.55
    col_w = 5.95
    gap = 0.35
    positions = (0.55, 0.55 + col_w + gap)

    for path, cap, left_in in zip((left, right), captions, positions):
        if path and path.is_file():
            _place_image_contain(
                slide, path, left_in, img_top, col_w - 0.2, max_h, with_frame=True
            )
        cap_y = img_top + max_h + 0.22
        cap_box = slide.shapes.add_textbox(
            Inches(left_in), Inches(cap_y), Inches(col_w - 0.2), Inches(0.32)
        )
        cap_box.text_frame.text = cap
        for p in cap_box.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                _set_font_safe(r, FONT_FALLBACK, 9, color=COLOR_TEXT)

    _add_footer(slide)
    set_notes(slide, spec.notes)
    return slide


def render_data_highlight(prs: Presentation, spec: DefenseSlide, slide_index: int):
    slide = _blank_slide(prs)
    _add_slide_number(slide, slide_index)
    _add_slide_title(slide, spec.title)

    hero = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.4), Inches(1.75), Inches(10.5), Inches(2.05)
    )
    hero.fill.solid()
    hero.fill.fore_color.rgb = COLOR_PRIMARY
    hero.line.fill.background()

    metric_box = slide.shapes.add_textbox(Inches(1.7), Inches(2.0), Inches(9.9), Inches(1.1))
    metric_box.text_frame.text = "~1.146 EUR/año"
    for p in metric_box.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            _set_font_safe(r, FONT_TITLE, 40, bold=True, color=COLOR_WHITE)

    sub = slide.shapes.add_textbox(Inches(1.7), Inches(3.15), Inches(9.9), Inches(0.4))
    sub.text_frame.text = "Coste total orientativo (infra cloud, dominio, correo)"
    for p in sub.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            _set_font_safe(r, FONT_BODY, 13, color=COLOR_WHITE)

    extra = [b for b in spec.bullets if "1.146" not in b]
    box = slide.shapes.add_textbox(Inches(0.62), Inches(4.2), Inches(12.0), Inches(2.5))
    _fill_bullets(box.text_frame, extra[:2], size=15)
    _add_footer(slide)
    set_notes(slide, spec.notes)
    return slide


def render_closing_statement(prs: Presentation, spec: DefenseSlide, slide_index: int):
    slide = _blank_slide(prs)
    _add_slide_number(slide, slide_index)
    _add_slide_title(slide, spec.title, size=20)

    stmt = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(1.35))
    stmt.text_frame.text = "¿Preguntas?"
    for p in stmt.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            _set_font_safe(r, FONT_TITLE, 44, bold=True, color=COLOR_ACCENT)

    checklist = [b for b in spec.bullets if "Preguntas" not in b and "?" not in b]
    box = slide.shapes.add_textbox(Inches(1.6), Inches(3.75), Inches(10.1), Inches(2.8))
    _fill_bullets(box.text_frame, checklist, size=16)
    for p in box.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER

    _add_footer(slide)
    set_notes(slide, spec.notes)
    return slide


def render_slide(
    prs: Presentation,
    spec: DefenseSlide,
    image_path: Path | None = None,
    image_right: Path | None = None,
    slide_index: int = 1,
):
    pattern = spec.visual_pattern
    if pattern == "title_slide":
        return render_title_slide(prs, spec, slide_index)
    if pattern == "timeline":
        return render_timeline_slide(prs, spec, slide_index)
    if pattern == "content_accent":
        return render_content_accent(prs, spec, slide_index)
    if pattern == "table_styled":
        return render_table_styled(prs, spec, slide_index)
    if pattern == "diagram_stack":
        return render_diagram_stack(prs, spec, image_path, slide_index)
    if pattern == "dual_images":
        return render_dual_images(prs, spec, image_path, image_right, slide_index)
    if pattern == "data_highlight":
        return render_data_highlight(prs, spec, slide_index)
    if pattern == "closing_statement":
        return render_closing_statement(prs, spec, slide_index)
    return render_content_accent(prs, spec, slide_index)
