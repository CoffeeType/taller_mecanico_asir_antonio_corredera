#!/usr/bin/env python3
"""Adjunta notas del orador del manifiesto a un PPTX exportado desde HTML."""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

from defense_slides_manifest import OUTPUT_PPTX_FROM_HTML, SLIDES


def repo_root() -> Path:
    return Path(__file__).resolve().parents[1]


def attach_notes(pptx_path: Path) -> None:
    prs = Presentation(str(pptx_path))
    if len(prs.slides) != len(SLIDES):
        print(
            f"AVISO: PPTX tiene {len(prs.slides)} slides, manifiesto {len(SLIDES)}. "
            "Se adjuntan notas por índice hasta el mínimo común."
        )
    n = min(len(prs.slides), len(SLIDES))
    for i in range(n):
        spec = SLIDES[i]
        notes = (spec.notes or "").strip()
        if not notes:
            continue
        slide = prs.slides[i]
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes
    prs.save(str(pptx_path))
    print(f"Notas adjuntadas en {n} diapositivas -> {pptx_path}")


def main() -> None:
    root = repo_root()
    pptx = root / "docs" / OUTPUT_PPTX_FROM_HTML
    if len(sys.argv) > 1:
        pptx = Path(sys.argv[1]).resolve()
    if not pptx.is_file():
        print(f"No encontrado: {pptx}")
        print("Genera antes: node tools/export_defense_html_to_pptx.mjs")
        raise SystemExit(1)
    attach_notes(pptx)


if __name__ == "__main__":
    main()
