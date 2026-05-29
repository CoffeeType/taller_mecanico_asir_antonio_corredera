#!/usr/bin/env python3
"""Comprueba que el PPTX *_html.pptx es el entregable raster (no el dom-to-pptx roto)."""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from defense_slides_manifest import OUTPUT_PPTX_FROM_HTML, SLIDES


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    pptx = root / "docs" / OUTPUT_PPTX_FROM_HTML
    if not pptx.is_file():
        print(f"FALTA: {pptx}", file=sys.stderr)
        raise SystemExit(1)

    kb = pptx.stat().st_size / 1024
    prs = Presentation(str(pptx))
    n = len(prs.slides)
    errors: list[str] = []

    if kb < 2000:
        errors.append(
            f"Tamaño sospechoso ({kb:.0f} KB). El raster válido suele ser >2 MB. "
            "¿Tienes abierto el PPTX viejo de dom-to-pptx (~69 KB)?"
        )
    if n != len(SLIDES):
        errors.append(f"Diapositivas: {n}, esperadas {len(SLIDES)}")

    for i, slide in enumerate(prs.slides, 1):
        pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        if len(pics) != 1:
            errors.append(f"Slide {i}: {len(pics)} imágenes (esperada 1)")
        elif pics[0].width < prs.slide_width * 0.95:
            errors.append(f"Slide {i}: imagen no cubre el lienzo")

    notes_ok = sum(
        1
        for slide in prs.slides
        if (slide.notes_slide.notes_text_frame.text or "").strip()
    )

    print(f"Archivo: {pptx.name}")
    print(f"Tamaño: {kb:.0f} KB")
    print(f"Diapositivas: {n}")
    print(f"Con notas del orador: {notes_ok}/{n}")
    print(f"Lienzo: {prs.slide_width / 914400:.3f}\" x {prs.slide_height / 914400:.3f}\"")

    if errors:
        print("\nFALLO:")
        for e in errors:
            print(f"  - {e}")
        raise SystemExit(1)

    print("\nOK — estructura del entregable raster correcta.")
    print("Abre el fichero en PowerPoint: cada slide debe ser una foto nítida del HTML.")
    print("Si necesitas editar textos en slide, usa el HTML o PFC_Defensa_...pptx (Corporate Blue).")


if __name__ == "__main__":
    main()
